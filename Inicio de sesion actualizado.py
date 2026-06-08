import tkinter as tk
from tkinter import messagebox
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import hashlib
import os
from datetime import datetime


EXCEL_FILE = "Prueba registro.xlsx"
SHEET_NAME = "Usuarios"
HEADERS    = ["ID", "Nombre", "Usuario", "Contraseña", "Fecha Registro"]


BG       = "#F7F8FA"
CARD     = "#FFFFFF"
ACCENT   = "#4F46E5"        
ACCENT_H = "#4338CA"
TEXT     = "#1E1E2D"
SUBTEXT  = "#6B7280"
BORDER   = "#E5E7EB"
ERROR    = "#EF4444"
SUCCESS  = "#10B981"
FONT_FAM = "Segoe UI"


def _init_excel():
    """Crea el archivo Excel con encabezados si no existe."""
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET_NAME

        
        header_fill = PatternFill("solid", fgColor="4F46E5")
        header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        center      = Alignment(horizontal="center", vertical="center")
        thin        = Side(style="thin", color="D1D5DB")
        border      = Border(left=thin, right=thin, top=thin, bottom=thin)

        col_widths = [8, 22, 20, 42, 22]
        for col_idx, (header, width) in enumerate(zip(HEADERS, col_widths), 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = center
            cell.border    = border
            ws.column_dimensions[cell.column_letter].width = width

        ws.row_dimensions[1].height = 28
        wb.save(EXCEL_FILE)


def _hash(password: str) -> str:
    return hashlib.sha256(password.encode()).hexdigest()


def _next_id() -> int:
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb[SHEET_NAME]
    return ws.max_row  


def _user_exists(usuario: str) -> bool:
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb[SHEET_NAME]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[2] and row[2].strip().lower() == usuario.strip().lower():
            return True
    return False


def _register_user(nombre: str, usuario: str, password: str):
    wb   = openpyxl.load_workbook(EXCEL_FILE)
    ws   = wb[SHEET_NAME]
    uid  = _next_id()
    fecha = datetime.now().strftime("%d/%m/%Y %H:%M")

    thin   = Side(style="thin", color="E5E7EB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    row_data = [uid, nombre.strip(), usuario.strip(), _hash(password), fecha]
    new_row  = ws.max_row + 1

    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=new_row, column=col_idx, value=value)
        cell.border    = border
        cell.alignment = center
        cell.font      = Font(name="Arial", size=10)
        if new_row % 2 == 0:
            cell.fill = PatternFill("solid", fgColor="F3F4F6")

    ws.row_dimensions[new_row].height = 22
    wb.save(EXCEL_FILE)


def _validate_login(usuario: str, password: str):
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb[SHEET_NAME]
    hashed = _hash(password)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if (row[2] and row[2].strip().lower() == usuario.strip().lower()
                and row[3] == hashed):
            return row[1]  
    return None



def _entry(parent, placeholder: str, show: str = "") -> tk.Entry:
    frame = tk.Frame(parent, bg=CARD, highlightbackground=BORDER,
                     highlightthickness=1, bd=0)
    frame.pack(fill="x", pady=6)

    e = tk.Entry(frame, font=(FONT_FAM, 11), bg=CARD, fg=SUBTEXT,
                 bd=0, relief="flat", show=show,
                 insertbackground=TEXT, highlightthickness=0)
    e.pack(padx=14, pady=10, fill="x")
    e.insert(0, placeholder)

    def on_focus_in(_):
        if e.get() == placeholder:
            e.delete(0, "end")
            e.config(fg=TEXT, show=show)
        frame.config(highlightbackground=ACCENT)

    def on_focus_out(_):
        if e.get() == "":
            e.insert(0, placeholder)
            e.config(fg=SUBTEXT, show="")
        frame.config(highlightbackground=BORDER)

    e.bind("<FocusIn>",  on_focus_in)
    e.bind("<FocusOut>", on_focus_out)
    e._placeholder = placeholder
    return e


def _get_val(entry: tk.Entry) -> str:
    val = entry.get()
    return "" if val == entry._placeholder else val


def _btn(parent, text: str, command, bg=ACCENT, fg="white") -> tk.Button:
    b = tk.Button(parent, text=text, command=command,
                  font=(FONT_FAM, 11, "bold"),
                  bg=bg, fg=fg, activebackground=ACCENT_H,
                  activeforeground="white", bd=0, relief="flat",
                  cursor="hand2", pady=11)
    b.pack(fill="x", pady=(10, 4))

    def on_enter(_): b.config(bg=ACCENT_H if bg == ACCENT else "#E5E7EB")
    def on_leave(_): b.config(bg=bg)
    b.bind("<Enter>", on_enter)
    b.bind("<Leave>", on_leave)
    return b


def _label(parent, text: str, size=11, color=TEXT, bold=False, pady=0):
    tk.Label(parent, text=text,
             font=(FONT_FAM, size, "bold" if bold else "normal"),
             bg=CARD, fg=color).pack(pady=pady)


def _card(root, width=380) -> tk.Frame:
    wrap = tk.Frame(root, bg=BG)
    wrap.pack(expand=True, fill="both")
    card = tk.Frame(wrap, bg=CARD, bd=0,
                    highlightbackground=BORDER, highlightthickness=1)
    card.place(relx=.5, rely=.5, anchor="center", width=width)
    return card



class LoginWindow:
    def __init__(self, root: tk.Tk):
        self.root = root
        root.title("Inicio de sesión")
        root.geometry("460x520")
        root.resizable(False, False)
        root.configure(bg=BG)
        self._build()

    def _build(self):
        for w in self.root.winfo_children():
            w.destroy()

        card = _card(self.root)

        inner = tk.Frame(card, bg=CARD)
        inner.pack(padx=38, pady=40, fill="x")

        
        tk.Label(inner, text="🔐", font=(FONT_FAM, 32), bg=CARD).pack(pady=(0, 4))
        _label(inner, "Bienvenido de nuevo", size=17, bold=True)
        _label(inner, "Ingresa tus credenciales para continuar",
               size=10, color=SUBTEXT, pady=(2, 18))

        self.e_user = _entry(inner, "Usuario")
        self.e_pass = _entry(inner, "Contraseña", show="•")

        self.msg = tk.Label(inner, text="", font=(FONT_FAM, 9),
                            bg=CARD, fg=ERROR)
        self.msg.pack()

        _btn(inner, "Iniciar sesión", self._login)

        sep = tk.Frame(inner, bg=CARD)
        sep.pack(fill="x", pady=(12, 0))
        tk.Label(sep, text="¿No tienes cuenta?", font=(FONT_FAM, 10),
                 bg=CARD, fg=SUBTEXT).pack(side="left")
        reg_btn = tk.Label(sep, text=" Regístrate", font=(FONT_FAM, 10, "bold"),
                           bg=CARD, fg=ACCENT, cursor="hand2")
        reg_btn.pack(side="left")
        reg_btn.bind("<Button-1>", lambda _: RegisterWindow(self.root))

    def _login(self):
        usuario  = _get_val(self.e_user)
        password = _get_val(self.e_pass)

        if not usuario or not password:
            self.msg.config(text="✖  Por favor completa todos los campos.")
            return

        nombre = _validate_login(usuario, password)
        if nombre:
            self._show_welcome(nombre)
        else:
            self.msg.config(text="✖  Usuario o contraseña incorrectos.")

    def _show_welcome(self, nombre: str):
        for w in self.root.winfo_children():
            w.destroy()

        card = _card(self.root)
        inner = tk.Frame(card, bg=CARD)
        inner.pack(padx=38, pady=50, fill="x")

        tk.Label(inner, text="✅", font=(FONT_FAM, 36), bg=CARD).pack(pady=(0, 8))
        _label(inner, f"¡Hola, {nombre}!", size=17, bold=True)
        _label(inner, "Has iniciado sesión correctamente.",
               size=10, color=SUBTEXT, pady=(4, 24))

        _btn(inner, "Cerrar sesión", self._build,
             bg="#F3F4F6", fg=TEXT)



class RegisterWindow:
    def __init__(self, master: tk.Tk):
        self.win = tk.Toplevel(master)
        self.win.title("Crear cuenta")
        self.win.geometry("460x580")
        self.win.resizable(False, False)
        self.win.configure(bg=BG)
        self.win.grab_set()
        self._build()

    def _build(self):
        card = _card(self.win)
        inner = tk.Frame(card, bg=CARD)
        inner.pack(padx=38, pady=36, fill="x")

        tk.Label(inner, text="👤", font=(FONT_FAM, 32), bg=CARD).pack(pady=(0, 4))
        _label(inner, "Crear cuenta", size=17, bold=True)
        _label(inner, "Completa los datos para registrarte",
               size=10, color=SUBTEXT, pady=(2, 18))

        self.e_nombre  = _entry(inner, "Nombre completo")
        self.e_user    = _entry(inner, "Usuario")
        self.e_pass    = _entry(inner, "Contraseña", show="•")
        self.e_confirm = _entry(inner, "Confirmar contraseña", show="•")

        self.msg = tk.Label(inner, text="", font=(FONT_FAM, 9),
                            bg=CARD, fg=ERROR, wraplength=300)
        self.msg.pack()

        _btn(inner, "Registrarse", self._register)

        lnk = tk.Label(inner, text="← Volver al inicio de sesión",
                       font=(FONT_FAM, 10), bg=CARD, fg=ACCENT, cursor="hand2")
        lnk.pack(pady=(10, 0))
        lnk.bind("<Button-1>", lambda _: self.win.destroy())

    def _register(self):
        nombre   = _get_val(self.e_nombre)
        usuario  = _get_val(self.e_user)
        password = _get_val(self.e_pass)
        confirm  = _get_val(self.e_confirm)

        if not all([nombre, usuario, password, confirm]):
            self.msg.config(text="✖  Por favor completa todos los campos.")
            return
        if len(password) < 6:
            self.msg.config(text="✖  La contraseña debe tener al menos 6 caracteres.")
            return
        if password != confirm:
            self.msg.config(text="✖  Las contraseñas no coinciden.")
            return
        if _user_exists(usuario):
            self.msg.config(text="✖  Ese nombre de usuario ya está en uso.")
            return

        _register_user(nombre, usuario, password)
        messagebox.showinfo(
            "¡Registro exitoso!",
            f"✅ Cuenta creada para {nombre}.\nYa puedes iniciar sesión.",
            parent=self.win
        )
        self.win.destroy()



if __name__ == "__main__":
    _init_excel()
    root = tk.Tk()
    root.update_idletasks()
    w, h = 460, 520
    x = (root.winfo_screenwidth()  - w) // 2
    y = (root.winfo_screenheight() - h) // 2
    root.geometry(f"{w}x{h}+{x}+{y}")
    LoginWindow(root)
    root.mainloop()
